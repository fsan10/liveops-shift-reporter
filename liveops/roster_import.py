"""Bounded spreadsheet parsing for additive roster imports; never evaluates formulas."""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO, StringIO
from typing import Any
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 2000
HEADERS = ["星期", "日期", "时段", "主播", "投手", "休息", "日常拍摄"]
ALIASES = {"主播姓名": "主播", "运营": "投手", "投手姓名": "投手", "拍摄": "日常拍摄"}


@dataclass
class ImportRow:
    number: int
    work_date: date | None = None
    start_time: str = ""
    end_time: str = ""
    host_name: str = ""
    promoter: str = ""
    rest: str = ""
    shoot: str = ""
    error: str = ""


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _text(value).replace("/", "-").replace("年", "-").replace("月", "-").rstrip("日")
    try:
        return datetime.strptime(raw, "%Y-%m-%d").date()
    except ValueError as exc:
        raise ValueError("日期请填写完整年月日，例如 2026-11-01。") from exc


def _time(value: Any) -> str:
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    raw = _text(value).replace("：", ":")
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(raw, fmt).strftime("%H:%M:%S")
        except ValueError:
            pass
    raise ValueError("时间请填写 HH:MM 或 HH:MM:SS。")


def _parse_rows(rows: Any) -> list[ImportRow]:
    columns: dict[str, int] | None = None
    previous_date: Any = None
    parsed: list[ImportRow] = []
    for number, values in enumerate(rows, 1):
        if number > MAX_ROWS + 1:
            raise ValueError(f"最多导入 {MAX_ROWS} 行，请拆分文件。")
        if not any(_text(value) for value in values):
            continue
        if columns is None:
            names = [ALIASES.get(_text(value), _text(value)) for value in values]
            if len([name for name in names if name]) != len({name for name in names if name}):
                raise ValueError("表头存在重复列，请使用导入模板。")
            columns = {name: i for i, name in enumerate(names) if name}
            if not {"日期", "主播"} <= columns.keys() or not (
                "时段" in columns or {"开始时间", "结束时间"} <= columns.keys()
            ):
                raise ValueError("缺少表头：必须包含日期、主播和时段（或开始时间、结束时间）。")
            continue

        def cell(name: str) -> Any:
            index = columns.get(name)
            return values[index] if index is not None and index < len(values) else None

        row = ImportRow(number=number)
        parsed.append(row)
        try:
            if _text(cell("日期")):
                previous_date = cell("日期")
            row.work_date = _date(previous_date)
            if any(_text(value).startswith("=") for value in values):
                raise ValueError("不支持公式，请将公式粘贴为值后导入。")
            if not 2000 <= row.work_date.year <= 2100:
                raise ValueError("排班年份须在 2000 至 2100 年之间。")
            row.host_name, row.promoter = _text(cell("主播")), _text(cell("投手"))
            row.rest, row.shoot = _text(cell("休息")), _text(cell("日常拍摄"))
            has_time = any(_text(cell(name)) for name in ("时段", "开始时间", "结束时间"))
            if not row.host_name and not has_time:
                if not row.rest and not row.shoot:
                    raise ValueError("请填写主播和时段，或休息/日常拍摄备注。")
            else:
                if not row.host_name:
                    raise ValueError("请填写主播姓名。")
                if "时段" in columns:
                    parts = re.split(r"\s*[-–—~～至]\s*", _text(cell("时段")))
                    if len(parts) != 2:
                        raise ValueError("时段请填写 07:00-11:30，跨午夜可填写 20:30-01:00。")
                    row.start_time, row.end_time = map(_time, parts)
                else:
                    row.start_time, row.end_time = _time(cell("开始时间")), _time(cell("结束时间"))
            if len(row.host_name) > 120 or len(row.promoter) > 120:
                raise ValueError("主播和投手姓名不能超过 120 字。")
            if len(row.rest) > 500 or len(row.shoot) > 500:
                raise ValueError("休息和拍摄备注不能超过 500 字。")
        except ValueError as exc:
            row.error = str(exc)
    if columns is None or not parsed:
        raise ValueError("文件没有可导入的排班行。")
    return parsed


def read_roster(content: bytes, filename: str) -> list[ImportRow]:
    if not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("请选择不超过 5 MB 的非空文件。")
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("gb18030")
            except UnicodeDecodeError as exc:
                raise ValueError("CSV 编码无法识别，请保存为 UTF-8 CSV。") from exc
        try:
            return _parse_rows(csv.reader(StringIO(text)))
        except csv.Error as exc:
            raise ValueError("CSV 内容无法读取，请检查列格式并重新保存。") from exc
    if suffix != "xlsx":
        raise ValueError("仅支持 .xlsx 和 .csv；旧版 .xls 请在 Excel 中另存为 .xlsx。")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 30 * 1024 * 1024:
                raise ValueError("Excel 解压后过大，请简化文件。")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            # First sheet only, independent of the operator's last active worksheet.
            sheet = workbook.worksheets[0]
            sheet.reset_dimensions()
            return _parse_rows(sheet.iter_rows(max_col=32, values_only=True))
        finally:
            workbook.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Excel 文件无法读取，请确认未加密且为有效的 .xlsx 文件。") from exc


def roster_template(today: date) -> bytes:
    """Generate the app's downloadable template (no local personnel or settings)."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "排班导入"
    sheet.append(HEADERS)
    sheet.append(["", today, "07:00-11:30", "示例主播A", "示例投手A", "", ""])
    sheet.append(["", today, "11:30-16:00", "示例主播B", "示例投手B", "", ""])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:G3"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="284B7E")
    for column, width in zip("ABCDEFG", (12, 16, 24, 20, 20, 28, 28), strict=True):
        sheet.column_dimensions[column].width = width
    for row in (2, 3):
        sheet.cell(row, 2).number_format = "yyyy-mm-dd"
    notes = workbook.create_sheet("填写说明")
    for line in (
        "每行一个班次；请替换示例姓名、日期和时段。只读取第一张工作表。",
        "必填列：日期、时段、主播。可选列：星期、投手、休息、日常拍摄。",
        "也可用开始时间、结束时间两列替代时段，支持 Excel 时间单元格。",
        "日期须含完整年月日。同一日期后续行可留空或合并日期单元格。",
        "跨午夜填写 20:30-01:00，归属开始日期；不支持公式，请粘贴为值。",
        "历史日期跳过；重复班次跳过；有错误时整批不写入。不会覆盖已有班次或非空备注。",
        "只有备注的日期可将主播、时段留空，填写休息或日常拍摄。",
        "导入后使用现有店铺和账户默认配置，班次默认启用；不会立即触发手动采集。",
        "支持 .xlsx 或本系统导出的 .csv，不超过 5 MB、2000 行。",
    ):
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
